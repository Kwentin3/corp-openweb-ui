"""Source-backed vertical worksheet composition, without routing cell data through an LLM.

OfficeCLI has no cross-workbook range relocation operation. This narrow adapter
uses openpyxl's reader/writer and formula tokenizer; it does not evaluate formulas.
Unsupported semantics fail before publishing. Original external links and cached
formula values are retained, not replaced by guesses or recalculated numbers.
"""
from __future__ import annotations

from copy import copy, deepcopy
from dataclasses import dataclass
from pathlib import Path
import posixpath
import re
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import quote_sheetname


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"s": MAIN}
CELL = re.compile(r"(\$?[A-Za-z]{1,3})(\$?)([1-9][0-9]*)\Z")
POSITION_FUNCTIONS = {"ROW(", "ROWS(", "COLUMN(", "COLUMNS(", "INDIRECT(", "ADDRESS(", "CELL(", "SHEET(", "SHEETS("}


class CompositionUnsupported(ValueError):
    """Cannot promise a faithful transformation for this source feature."""


@dataclass(frozen=True)
class Source:
    path: Path
    target_sheet: str


def _unquote_sheet(value: str) -> str:
    return value[1:-1].replace("''", "'") if value.startswith("'") and value.endswith("'") else value


def _reference(value: str, current_sheet: str) -> tuple[str, str]:
    if "!" in value:
        sheet, address = value.rsplit("!", 1)
        return _unquote_sheet(sheet), address
    return current_sheet, value


def _shift_address(address: str, offset: int) -> str:
    parts = address.split(":")
    if len(parts) > 2:
        raise CompositionUnsupported("3D and non-A1 references need a different composition route")
    shifted = []
    for part in parts:
        match = CELL.fullmatch(part)
        if not match:
            raise CompositionUnsupported("Named, structured, and full-column/row references are not supported for stacking")
        row = int(match[3]) + offset
        if row > 1048576:
            raise CompositionUnsupported("The composed worksheet exceeds Excel's row limit")
        shifted.append(f"{match[1]}{match[2]}{row}")
    return ":".join(shifted)


def relocate_formula(formula: str, current_sheet: str, target: str,
                     offsets: dict[str, int], external_base: int) -> str:
    tokens = Tokenizer(formula).items
    for token in tokens:
        if token.type == "FUNC" and token.value.upper() in POSITION_FUNCTIONS:
            raise CompositionUnsupported("Position-dependent formulas cannot be stacked without changing their meaning")
        if token.subtype != "RANGE":
            continue
        sheet, address = _reference(token.value, current_sheet)
        external = re.match(r"^\[([1-9][0-9]*)\](.*)$", sheet)
        if external:
            token.value = f"{quote_sheetname(f'[{int(external[1]) + external_base}]{external[2]}')}!{address}"
            continue
        if sheet not in offsets:
            raise CompositionUnsupported("A formula refers to an unavailable sheet or unsupported external reference")
        token.value = f"{quote_sheetname(target)}!{_shift_address(address, offsets[sheet])}"
    return "=" + "".join(token.value for token in tokens)


def _sheet_xml(archive: ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
    return {s.get("name"): posixpath.normpath(posixpath.join("xl", rels[s.get(f"{{{REL}}}id")])).lstrip("/")
            for s in workbook.find(f"{{{MAIN}}}sheets")}


def _check_package(path: Path) -> int:
    with ZipFile(path) as archive:
        expanded_size = sum(i.file_size for i in archive.infolist())
        if expanded_size > 128 * 1024 * 1024:
            raise CompositionUnsupported("Source workbook expands beyond the 128 MiB processing limit")
        for name in archive.namelist():
            if any(name.startswith(prefix) for prefix in ("xl/pivot", "xl/slicer", "xl/embeddings/", "xl/activeX/", "xl/charts/")) or name.endswith("vbaProject.bin"):
                raise CompositionUnsupported("Pivot tables, charts, embedded objects, and macros are not supported for stacking")
            if name.startswith("xl/drawings/drawing") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                if any(ET.QName(n).localname in {"sp", "grpSp", "cxnSp", "graphicFrame"} for n in root.iter()):
                    raise CompositionUnsupported("Non-picture drawing objects need a different composition route")
        return expanded_size


def _offsets(workbook) -> dict[str, int]:
    sizes = {s.title: s.max_row for s in workbook}
    for sheet in workbook:
        if sheet.tables or sheet.conditional_formatting or sheet.data_validations.dataValidation:
            raise CompositionUnsupported("Tables, conditional formatting, and data validation are not supported for stacking")
        if sheet.max_row * sheet.max_column > 500000:
            raise CompositionUnsupported("A source sheet exceeds the 500000-cell processing limit")
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                if not isinstance(cell.value, str):
                    raise CompositionUnsupported("Array and data-table formulas are not supported for stacking")
                for token in Tokenizer(cell.value).items:
                    if token.subtype == "RANGE":
                        name, address = _reference(token.value, sheet.title)
                        if name in sizes:
                            _shift_address(address, 0)  # fail closed before allocating target rows
                            sizes[name] = max(sizes[name], *(int(CELL.fullmatch(a)[3]) for a in address.split(":")))
    offset = 1  # first row of each block labels the original sheet
    result = {}
    for name, size in sizes.items():
        result[name] = offset
        offset += size + 2
    if offset > 1048576:
        raise CompositionUnsupported("The composed worksheet exceeds Excel's row limit")
    return result


def _copy_cell(source, target, formula: str | None) -> None:
    target.value = formula if formula is not None else copy(source.value)
    target.data_type = source.data_type
    _copy_style(source, target)
    if source.comment:
        target.comment = copy(source.comment)
    if source.hyperlink:
        # Internal hyperlinks are address-bearing objects too; do not silently leave them stale.
        if source.hyperlink.location:
            raise CompositionUnsupported("Internal hyperlinks need a different composition route")
        target.hyperlink = copy(source.hyperlink)
        target.hyperlink.ref = target.coordinate


def _copy_style(source, target) -> None:
    for attribute in ("font", "fill", "border", "alignment", "protection"):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.number_format = source.number_format


def _copy_dimension(dimension, target):
    cloned = copy(dimension)
    cloned.parent = target
    cloned._style = None
    _copy_style(dimension, cloned)
    return cloned


def _restore_caches(result: Path, caches: dict[str, dict[str, tuple[str | None, bytes | None]]]) -> None:
    """Keep original numeric precision and last-calculated formula values.

    openpyxl deliberately doesn't write formula caches. Relocation preserves the
    calculation, so retain the original cache (including existing errors/absence).
    Numeric literals also retain their original XML representation rather than
    going through Python's float serialization. Date epochs must match. No
    evaluation or fabrication happens here. Excel still recalculates on open.
    """
    temporary = result.with_suffix(".cached.xlsx")
    with ZipFile(result) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        paths = {path: caches[name] for name, path in _sheet_xml(source).items()}
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in paths:
                root = ET.fromstring(data)
                for cell in root.findall(".//s:sheetData/s:row/s:c", NS):
                    cached = paths[info.filename].get(cell.get("r"))
                    if cached is None:
                        continue
                    kind, value = cached
                    if kind:
                        cell.set("t", kind)
                    else:
                        cell.attrib.pop("t", None)
                    for old in cell.findall("s:v", NS):
                        cell.remove(old)
                    if value is not None:
                        cell.append(ET.fromstring(value))
                data = ET.tostring(root)
            target.writestr(info, data)
    temporary.replace(result)


def compose(sources: list[Source], result: Path) -> dict:
    if not sources or len(sources) > 64:
        raise CompositionUnsupported("Provide between 1 and 64 source workbooks")
    titles = [s.target_sheet for s in sources]
    if len({n.casefold() for n in titles}) != len(titles) or any(not n or len(n) > 31 or re.search(r"[\\/*?:\[\]]", n) for n in titles):
        raise CompositionUnsupported("Target sheet names must be unique valid Excel names of at most 31 characters")
    output = Workbook()
    output.remove(output.active)
    caches = {}
    receipt = {"source_workbooks": len(sources), "source_sheets": 0, "copied_cells": 0,
               "formulas": 0, "pictures": 0, "external_links": 0, "source_error_cells": 0,
               "layout": "source sheets stacked in original order with labelled blocks; column widths use the widest source column"}
    expanded_bytes = sum(_check_package(item.path) for item in sources)
    if expanded_bytes > 128 * 1024 * 1024:
        raise CompositionUnsupported("Combined workbooks expand beyond the 128 MiB processing limit")
    for item in sources:
        workbook = load_workbook(item.path, rich_text=True)
        if item is sources[0]:
            output.epoch = workbook.epoch
            output.loaded_theme = workbook.loaded_theme
        elif workbook.epoch != output.epoch:
            raise CompositionUnsupported("Workbooks with different date epochs require explicit date conversion")
        elif workbook.loaded_theme != output.loaded_theme:
            raise CompositionUnsupported("Workbooks with different themes require explicit color and font conversion")
        if workbook.defined_names:
            raise CompositionUnsupported("Workbook named ranges need a different composition route")
        offsets = _offsets(workbook)
        target = output.create_sheet(item.target_sheet)
        external_base = len(output._external_links)
        output._external_links.extend(deepcopy(workbook._external_links))
        receipt["external_links"] += len(workbook._external_links)
        caches[item.target_sheet] = {}
        with ZipFile(item.path) as archive:
            for name, path in _sheet_xml(archive).items():
                for cell in ET.fromstring(archive.read(path)).findall(".//s:sheetData/s:row/s:c", NS):
                    receipt["source_error_cells"] += int(cell.get("t") == "e")
                    if cell.find("s:f", NS) is not None or cell.get("t") in {None, "n", "b", "e"}:
                        value = cell.find("s:v", NS)
                        caches[item.target_sheet][_shift_address(cell.get("r"), offsets[name])] = (
                            cell.get("t"), ET.tostring(value) if value is not None else None)
        for sheet in workbook:
            offset = offsets[sheet.title]
            target.cell(offset, 1, sheet.title)
            receipt["source_sheets"] += 1
            for row in sheet:
                for cell in row:
                    if isinstance(cell, MergedCell) or (cell.value is None and not cell.has_style):
                        continue
                    formula = relocate_formula(cell.value, sheet.title, target.title, offsets, external_base) if cell.data_type == "f" else None
                    _copy_cell(cell, target.cell(cell.row + offset, cell.column), formula)
                    receipt["copied_cells"] += 1
                    if receipt["copied_cells"] > 250000:
                        raise CompositionUnsupported("Combined workbooks exceed the 250000-cell processing limit")
                    receipt["formulas"] += int(formula is not None)
            for merged in sheet.merged_cells.ranges:
                target.merge_cells(_shift_address(str(merged), offset))
            for index, dimension in sheet.row_dimensions.items():
                target.row_dimensions[index + offset] = _copy_dimension(dimension, target)
                target.row_dimensions[index + offset].index = index + offset
            for key, dimension in sheet.column_dimensions.items():
                existing = target.column_dimensions.get(key)
                if existing is None or (dimension.width or 0) > (existing.width or 0):
                    target.column_dimensions[key] = _copy_dimension(dimension, target)
            for picture in sheet._images:
                anchor = deepcopy(picture.anchor)
                if not hasattr(anchor, "_from"):
                    raise CompositionUnsupported("Absolute picture anchors need a different composition route")
                anchor._from.row += offset
                if hasattr(anchor, "to"):
                    anchor.to.row += offset
                geometry = anchor.pic.spPr.prstGeom if anchor.pic is not None else None
                if geometry is not None and geometry.avLst is not None:
                    if geometry.avLst.gd is not None:
                        raise CompositionUnsupported("Adjusted picture geometry needs a different composition route")
                    # openpyxl 3.1.5 emits an empty avLst in the inherited xdr namespace.
                    # Omit the optional empty list; preserve the preset geometry itself.
                    geometry.avLst = None
                if anchor.pic is not None:
                    anchor.pic.nvPicPr.cNvPr.id = len(target._images) + 1
                picture.anchor = anchor
                target.add_image(picture)
                receipt["pictures"] += 1
        workbook.close()
    output.save(result)
    _restore_caches(result, caches)
    return receipt


if __name__ == "__main__":
    import json
    import sys
    from .xlsx_verification import verify, CompositionVerificationError

    try:
        if sys.platform != "win32":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
        plan = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
        sources = [Source(Path(item["path"]), item["target_sheet"]) for item in plan]
        destination = Path(sys.argv[2])
        receipt = compose(sources, destination)
        receipt["verified"] = verify([source.path for source in sources], destination)
        print(json.dumps(receipt, ensure_ascii=False))
    except (CompositionUnsupported, CompositionVerificationError) as error:
        print(json.dumps({"reason": "composition_not_supported_or_not_verified", "detail": str(error)}))
        sys.exit(2)
    except Exception as error:
        # Do not leak workbook contents, original paths, or library tracebacks to the model.
        print(json.dumps({"reason": "composition_failed", "error_type": type(error).__name__}))
        sys.exit(2)
