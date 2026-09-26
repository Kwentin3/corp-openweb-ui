"""Independent source/result comparison. Private workbooks stay outside Git.

Run with a directory of source-*.xlsx and the produced combined workbook.
This checks formula dependency addresses in reverse (output -> source), not by
calling the production formula relocation routine again.
"""
from bisect import bisect_right
from copy import copy
from hashlib import sha256
import re
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

class CompositionVerificationError(ValueError):
    pass


def formula_tokens(value, current, source_names, boundaries, external_base, reverse):
    result = []
    for token in Tokenizer(value).items:
        if token.subtype != 'RANGE':
            result.append((token.type, token.subtype, token.value))
            continue
        sheet, sep, address = token.value.rpartition('!')
        if not sep:
            sheet, address = (current, token.value)
        sheet = sheet.strip("'").replace("''", "'")
        external = re.fullmatch('\\[([0-9]+)\\](.*)', sheet)
        if external:
            result.append(('external', int(external[1]) - (external_base if reverse else 0), external[2], address))
            continue
        if reverse and sheet != current:
            raise CompositionVerificationError('Formula unexpectedly points to a different destination sheet')
        refs = []
        for cell in address.split(':'):
            match = re.fullmatch('(\\$?[A-Za-z]+)(\\$?)([0-9]+)', cell)
            if not match:
                raise CompositionVerificationError('unexpected reference')
            row = int(match[3])
            if reverse:
                block = bisect_right(boundaries, row - 1) - 1
                if not block >= 0:
                    raise CompositionVerificationError('Source/result verification failed')
                name = source_names[block]
                row -= boundaries[block]
            else:
                name = sheet
            refs.append((name, match[1], match[2], row))
        result.append(('internal', refs))
    return result

def verify(paths, result):
    out = load_workbook(result, rich_text=True)
    out_values = load_workbook(result, data_only=True)
    if not len(out.worksheets) == len(paths):
        raise CompositionVerificationError('Source/result verification failed')
    counts = {'cells': 0, 'formulas': 0, 'cached_values': 0, 'pictures': 0, 'merged_ranges': 0}
    external_base = 0
    for index, path in enumerate(paths):
        src = load_workbook(path, rich_text=True)
        src_values = load_workbook(path, data_only=True)
        target = out.worksheets[index]
        target_values = out_values.worksheets[index]
        boundaries = []
        last = 0
        for sheet in src:
            hits = [r for r in range(last + 1, target.max_row + 1) if target.cell(r, 1).value == sheet.title]
            if not hits:
                raise CompositionVerificationError('source-sheet label missing')
            last = hits[0]
            boundaries.append(last)
        names = src.sheetnames
        for si, sheet in enumerate(src):
            offset = boundaries[si]
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    actual = target.cell(cell.row + offset, cell.column)
                    counts['cells'] += 1
                    if cell.data_type == 'f':
                        if not actual.data_type == 'f':
                            raise CompositionVerificationError('formula replaced by a constant')
                        if not formula_tokens(cell.value, sheet.title, names, boundaries, external_base, False) == formula_tokens(actual.value, target.title, names, boundaries, external_base, True):
                            raise CompositionVerificationError(f'formula dependency changed: workbook {index + 1} sheet {si + 1} cell {cell.coordinate}')
                        counts['formulas'] += 1
                    elif not cell.value == actual.value:
                        raise CompositionVerificationError(f'value changed: workbook {index + 1} sheet {si + 1} cell {cell.coordinate}')
                    if not src_values[sheet.title][cell.coordinate].value == target_values.cell(cell.row + offset, cell.column).value:
                        raise CompositionVerificationError(f'cached value changed: workbook {index + 1} sheet {si + 1} cell {cell.coordinate}')
                    counts['cached_values'] += 1
                    for attribute in ('font', 'fill', 'alignment', 'number_format', 'protection'):
                        if not copy(getattr(cell, attribute)) == copy(getattr(actual, attribute)):
                            raise CompositionVerificationError(f'cell format changed: {attribute}')
            for merged in sheet.merged_cells.ranges:
                shifted = str(merged).split(':')
                expected = ':'.join((re.sub('[0-9]+', lambda m: str(int(m[0]) + offset), s) for s in shifted))
                if not expected in target.merged_cells:
                    raise CompositionVerificationError('Source/result verification failed')
                counts['merged_ranges'] += 1
            counts['pictures'] += len(sheet._images)
        if not len(target._images) == sum((len(s._images) for s in src)):
            raise CompositionVerificationError('Source/result verification failed')
        source_images = [picture for sheet in src for picture in sheet._images]
        if [sha256(p._data()).hexdigest() for p in source_images] != [sha256(p._data()).hexdigest() for p in target._images]:
            raise CompositionVerificationError('Picture bytes changed')
        for a, b in zip(src._external_links, out._external_links[external_base:]):
            if not a.file_link.Target == b.file_link.Target:
                raise CompositionVerificationError('Source/result verification failed')
            if a.externalBook.sheetDataSet != b.externalBook.sheetDataSet:
                raise CompositionVerificationError('External dependency cached data changed')
        external_base += len(src._external_links)
    if not len(out._external_links) == external_base:
        raise CompositionVerificationError('Source/result verification failed')
    return counts
