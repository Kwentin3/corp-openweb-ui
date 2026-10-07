"""Real pinned CLI + adapter qualification, with synthetic files and local file transport.

Run inside the built adapter image: python tests/verify_capability_discovery.py.
This is not a provider or authenticated OpenWebUI end-to-end test.
"""

from hashlib import sha256
from io import BytesIO
import importlib.metadata
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile
from xml.etree import ElementTree

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image

from officecli_openapi_proof.app import create_app
from officecli_openapi_proof.config import load_settings
from officecli_openapi_proof.officecli import SubprocessOfficeCliExecutor


class LocalFileTransport:
    def __init__(self, sources):
        self.sources = sources
        self.result = None
        self.attached = []

    def verify_session(self, authorization):
        assert authorization == "Bearer synthetic-session"

    def has_nearest_pptx_attachment(self, chat_id, message_id, authorization):
        return False

    def download(self, file_id, authorization, destination):
        destination.write_bytes(self.sources[file_id].read_bytes())

    def upload(self, source, output_name, authorization, content_type=None):
        self.result = source.read_bytes()
        return {"id": "synthetic-result", "filename": output_name}

    def attach(self, chat_id, message_id, native_file, authorization, fallback_name=None, content_type=None):
        self.attached.append(native_file["id"])


class DiscoveryQualification(unittest.TestCase):
    def test_installed_dependencies_match_production_lock(self):
        for line in Path("/app/requirements.lock").read_text().splitlines():
            if line and not line.startswith("#"):
                name, version = line.split("==")
                self.assertEqual(importlib.metadata.version(name), version, name)

    def test_large_real_workbook_structure_is_complete_and_readonly(self):
        settings = load_settings()
        executor = SubprocessOfficeCliExecutor(settings)
        with TemporaryDirectory() as directory:
            source = Path(directory) / "many-sheets.xlsx"
            workbook = Workbook()
            names = [f"Actual sheet {i:03}" for i in range(137)]
            for index, name in enumerate(names):
                sheet = workbook.active if index == 0 else workbook.create_sheet()
                sheet.title = name
                sheet["A1"] = name
            workbook.save(source)
            before = source.read_bytes()
            files = LocalFileTransport({"source": source})
            client = TestClient(create_app(executor, files, settings))
            headers = {"Authorization": "Bearer synthetic-session", "X-OpenWebUI-Chat-Id": "chat",
                       "X-OpenWebUI-Message-Id": "message"}
            offset, recovered = 0, []
            while offset is not None:
                response = client.post("/v1/officecli/spreadsheets/inspect", headers=headers,
                    json={"file_id": "source", "command_payload": {"mode": "outline", "offset": offset}})
                self.assertEqual(response.status_code, 200, response.text)
                result = response.json()["officecli_result"]
                recovered.extend(sheet["name"] for sheet in result["data"]["sheets"])
                offset = result["pagination"]["next_offset"]
            self.assertEqual(recovered, names)
            self.assertEqual(source.read_bytes(), before)
            self.assertIsNone(files.result)
            self.assertEqual(files.attached, [])

    def test_official_help_and_objects_in_all_formats(self):
        settings = load_settings()
        executor = SubprocessOfficeCliExecutor(settings)
        self.assertEqual(executor.run("--version").text.strip(), settings.expected_version)
        headers = {"Authorization": "Bearer synthetic-session",
                   "X-OpenWebUI-Chat-Id": "synthetic-chat", "X-OpenWebUI-Message-Id": "synthetic-message"}
        with TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "image.png"
            Image.new("RGB", (16, 16), "blue").save(image)
            second_image = root / "second-image.png"
            Image.new("RGB", (16, 16), "green").save(second_image)
            sources = {fmt: root / f"source.{fmt}" for fmt in ("xlsx", "docx", "pptx")}
            workbook = Workbook()
            for index, name in enumerate(("Empty", "Actual A", "Actual B")):
                sheet = workbook.active if index == 0 else workbook.create_sheet()
                sheet.title = name
                sheet["A1"] = f"Keep {name}"
                sheet["B1"] = "=1+2"
                for number in range(index * 3):
                    sheet.add_image(SpreadsheetImage(image), f"D{number + 1}")
            workbook.save(sources["xlsx"])
            for fmt, commands in (
                ("docx", [{"command": "add", "parent": "/body", "type": "paragraph", "props": {"text": "Keep document text"}},
                          {"command": "add", "parent": "/body/p[1]", "type": "picture", "props": {"src": str(image)}}]),
                ("pptx", [{"command": "add", "parent": "/", "type": "slide", "props": {"layout": "blank"}},
                          {"command": "add", "parent": "/slide[1]", "type": "picture", "props": {"src": str(image), "x": "1cm", "y": "1cm", "width": "2cm", "height": "2cm"}},
                          {"command": "add", "parent": "/slide[1]", "type": "shape", "props": {"text": "Original title"}},
                          {"command": "add", "parent": "/", "type": "slide", "props": {"layout": "blank"}},
                          {"command": "add", "parent": "/slide[2]", "type": "picture", "props": {"src": str(second_image), "x": "1cm", "y": "1cm", "width": "2cm", "height": "2cm"}},
                          {"command": "add", "parent": "/slide[2]", "type": "shape", "props": {"text": "Keep editable second slide"}},
                          {"command": "add", "parent": "/slide[2]", "type": "table",
                           "props": {"data": "Item,Price;Template item,100", "x": "1cm", "y": "4cm", "width": "6cm", "height": "2cm"}}]),
            ):
                executor.run("create", str(sources[fmt]), "--json")
                executor.run("batch", str(sources[fmt]), "--stop-on-error", "--json", input_text=json.dumps(commands))
            files = LocalFileTransport(sources)
            client = TestClient(create_app(executor, files, settings))
            workflow = client.post("/v1/officecli/help", headers=headers, json={})
            self.assertEqual(workflow.status_code, 200, workflow.text)
            native_tools = json.loads(executor.run("mcp", input_text=json.dumps({
                "jsonrpc": "2.0", "id": 1, "method": "tools/list", "params": {}}) + "\n").text)
            official = next(t for t in native_tools["result"]["tools"] if t["name"] == "officecli")
            self.assertEqual(workflow.json()["content"], official["description"])
            schema = client.get("/openapi.json").json()
            description = schema["paths"]["/v1/officecli/help"]["post"]["description"]
            self.assertIn(official["description"], description)
            self.assertEqual(description.count(official["description"]), 1)
            catalog = client.post("/v1/officecli/skills/load", headers=headers, json={})
            self.assertEqual(catalog.status_code, 200, catalog.text)
            self.assertIn("word-form", catalog.json()["content"])
            self.assertIn("financial-model", catalog.json()["content"])
            for skill in ("word", "excel", "pptx", "word-form"):
                guide = client.post("/v1/officecli/skills/load", headers=headers, json={"skill": skill})
                self.assertEqual(guide.status_code, 200, guide.text)
                self.assertEqual(guide.json()["content"], executor.run("load_skill", skill).text)
            for fmt, kind, count in (("xlsx", "spreadsheets", 9), ("docx", "documents", 1), ("pptx", "presentations", 2)):
                with self.subTest(format=fmt):
                    before = sha256(sources[fmt].read_bytes()).hexdigest()
                    for topic in (fmt, f"{fmt} picture", f"{fmt} remove picture", "query", "get"):
                        response = client.post("/v1/officecli/help", headers=headers, json={"topic": topic})
                        self.assertEqual(response.status_code, 200, response.text)
                        self.assertTrue(response.json()["content"])
                    paths = []
                    offset = 0
                    while offset is not None:
                        response = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                            json={"file_id": fmt, "command_payload": {"command": "query", "selector": "picture", "limit": 2, "offset": offset}})
                        self.assertEqual(response.status_code, 200, response.text)
                        result = response.json()["officecli_result"]
                        self.assertEqual(result["data"]["matches"], count)
                        paths.extend(node["path"] for node in result["data"]["results"])
                        offset = result["pagination"]["next_offset"]
                    self.assertEqual(len(paths), count)
                    response = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                        json={"file_id": fmt, "command_payload": {"command": "get", "path": paths[0], "depth": 0}})
                    self.assertEqual(response.status_code, 200, response.text)
                    self.assertEqual(response.json()["officecli_result"]["data"]["results"][0]["path"], paths[0])
                    self.assertIsNone(files.result)
                    self.assertEqual(files.attached, [])
                    html = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                        json={"file_id": fmt, "command_payload": {"command": "view", "mode": "html"}})
                    self.assertEqual(html.status_code, 200, html.text[:300])
                    html_result = html.json()["officecli_result"]
                    if html_result.get("encoding") != "json-fragment":
                        self.assertEqual(html_result["format"], "html")
                        self.assertIn("html", html_result["content"].lower())
                    if fmt == "docx":
                        forms = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                            json={"file_id": fmt, "command_payload": {"command": "view", "mode": "forms"}})
                        self.assertEqual(forms.status_code, 200, forms.text)
                    raw_path = {"xlsx": "/workbook", "docx": "/document", "pptx": "/presentation"}[fmt]
                    raw = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                        json={"file_id": fmt, "command_payload": {"command": "raw", "path": raw_path}})
                    self.assertEqual(raw.status_code, 200, raw.text)
                    self.assertTrue(raw.json()["officecli_result"]["success"])
                    for payload in ({"command": "view", "mode": "issues"}, {"command": "validate"}):
                        check = client.post(f"/v1/officecli/{kind}/inspect", headers=headers,
                            json={"file_id": fmt, "command_payload": payload})
                        self.assertEqual(check.status_code, 200, check.text)
                        self.assertTrue(check.json()["officecli_result"]["success"])
                    screenshot = client.post("/v1/officecli/render", headers=headers,
                        json={"file_id": fmt, "format": fmt})
                    self.assertEqual(screenshot.status_code, 200, screenshot.text[:300] if screenshot.status_code != 200 else "")
                    with Image.open(BytesIO(screenshot.content)) as rendered:
                        self.assertEqual(rendered.format, "PNG")
                        self.assertGreater(rendered.width, 100)
                        self.assertGreater(rendered.height, 100)
                        self.assertGreater(len(rendered.convert("RGB").resize((100, 100)).getcolors(10001)), 2)
                    if fmt == "xlsx":
                        region = client.post("/v1/officecli/render", headers=headers,
                            json={"file_id": fmt, "format": fmt, "range": "Actual B!A1:B1"})
                        self.assertEqual(region.status_code, 200, region.text[:300] if region.status_code != 200 else "")
                        with Image.open(BytesIO(region.content)) as cropped, Image.open(BytesIO(screenshot.content)) as full:
                            self.assertLess(cropped.width * cropped.height, full.width * full.height)
                        response = client.post(f"/v1/officecli/{kind}/apply-batch", headers=headers,
                            json={"file_id": fmt, "output_name": "without-pictures.xlsx",
                                "commands": [{"command": "remove", "path": path} for path in reversed(paths)]})
                        self.assertEqual(response.status_code, 200, response.text)
                        self.assertEqual(response.json()["result_file_id"], "synthetic-result")
                        self.assertEqual(response.json()["download_url"], "/api/v1/files/synthetic-result/content")
                        self.assertEqual(files.attached, ["synthetic-result"])
                        final = root / "result.xlsx"
                        final.write_bytes(files.result)
                        result = json.loads(executor.run("query", str(final), "picture", "--json").text)
                        self.assertEqual(result["data"]["matches"], 0)
                        with ZipFile(final) as archive:
                            self.assertFalse(any(name.startswith("xl/media/") for name in archive.namelist()))
                        updated = load_workbook(final)
                        self.assertEqual(updated.sheetnames, ["Empty", "Actual A", "Actual B"])
                        self.assertEqual([sheet["A1"].value for sheet in updated], ["Keep Empty", "Keep Actual A", "Keep Actual B"])
                        self.assertEqual([sheet["B1"].value for sheet in updated], ["=1+2"] * 3)
                        files.result = None
                        files.attached.clear()
                    if fmt == "pptx":
                        response = client.post("/v1/officecli/presentations/apply-batch", headers=headers,
                            json={"file_id": fmt, "output_name": "edited-template.pptx",
                                "commands": [{"command": "set", "path": "/slide[1]/shape[1]",
                                              "props": {"text": "Updated title"}},
                                             {"command": "add", "parent": "/", "from": "/slide[2]"}]})
                        self.assertEqual(response.status_code, 200, response.text)
                        self.assertTrue(response.json()["source_bytes_preserved"])
                        self.assertEqual(files.attached, ["synthetic-result"])
                        final = root / "edited-template.pptx"
                        final.write_bytes(files.result)
                        def slide_objects(archive, slide_number):
                            slide = ElementTree.fromstring(archive.read(f"ppt/slides/slide{slide_number}.xml"))
                            presentation = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
                            drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                            return ({kind: len(slide.findall(f".//{presentation}{kind}"))
                                     for kind in ("sp", "pic", "graphicFrame")},
                                    [node.text for node in slide.findall(f".//{drawing}t")])
                        with ZipFile(sources[fmt]) as original, ZipFile(final) as edited:
                            original_media = {name: original.read(name) for name in original.namelist()
                                              if name.startswith("ppt/media/")}
                            edited_media = {name: edited.read(name) for name in edited.namelist()
                                            if name.startswith("ppt/media/")}
                            self.assertEqual(edited_media, original_media)
                            self.assertIn(b"Updated title", edited.read("ppt/slides/slide1.xml"))
                            self.assertEqual(slide_objects(edited, 2), slide_objects(original, 2))
                            self.assertEqual(slide_objects(edited, 3), slide_objects(original, 2))
                            self.assertIn(b"Keep editable second slide", edited.read("ppt/slides/slide3.xml"))
                        source_second = root / "source-template-slide.png"
                        second = root / "template-slide.png"
                        cloned = root / "cloned-slide.png"
                        executor.run("view", str(sources[fmt]), "screenshot", "--page", "2", "--render", "html", "--out", str(source_second))
                        executor.run("view", str(final), "screenshot", "--page", "2", "--render", "html", "--out", str(second))
                        executor.run("view", str(final), "screenshot", "--page", "3", "--render", "html", "--out", str(cloned))
                        self.assertEqual(source_second.read_bytes(), second.read_bytes())
                        self.assertEqual(second.read_bytes(), cloned.read_bytes())
                        files.sources[response.json()["result_file_id"]] = final
                        followup = client.post("/v1/officecli/presentations/apply-batch", headers=headers,
                            json={"file_id": response.json()["result_file_id"],
                                  "output_name": "edited-variant.pptx",
                                  "commands": [{"command": "set", "path": "/slide[3]/shape[1]",
                                                "props": {"text": "Visible variant"}},
                                               {"command": "set", "path": "/slide[3]/table[1]/tr[2]/tc[1]",
                                                "props": {"text": "Variant item"}}]})
                        self.assertEqual(followup.status_code, 200, followup.text)
                        self.assertEqual(followup.json()["source_file_id"], response.json()["result_file_id"])
                        self.assertEqual(followup.json()["source_sha256"], sha256(final.read_bytes()).hexdigest())
                        variant = root / "edited-variant.pptx"
                        variant.write_bytes(files.result)
                        with ZipFile(final) as previous, ZipFile(variant) as current:
                            self.assertEqual(slide_objects(current, 2), slide_objects(previous, 2))
                            self.assertIn(b"Visible variant", current.read("ppt/slides/slide3.xml"))
                            self.assertIn(b"Variant item", current.read("ppt/slides/slide3.xml"))
                            self.assertEqual(
                                {name: current.read(name) for name in current.namelist() if name.startswith("ppt/media/")},
                                {name: previous.read(name) for name in previous.namelist() if name.startswith("ppt/media/")},
                            )
                        changed = root / "visible-variant.png"
                        executor.run("view", str(variant), "screenshot", "--page", "3", "--render", "html", "--out", str(changed))
                        self.assertNotEqual(cloned.read_bytes(), changed.read_bytes())
                        files.result = None
                        files.attached.clear()
                    self.assertEqual(sha256(sources[fmt].read_bytes()).hexdigest(), before)

    def test_native_pptx_validation_owns_rejected_commands(self):
        settings = load_settings()
        executor = SubprocessOfficeCliExecutor(settings)
        files = LocalFileTransport({})
        client = TestClient(create_app(executor, files, settings))
        headers = {"Authorization": "Bearer synthetic-session", "X-OpenWebUI-Chat-Id": "chat",
                   "X-OpenWebUI-Message-Id": "message"}
        for commands in (
            [{"command": "set", "path": "/slide[1]/table[1]", "props": {"r1c1": "Budget"}}],
            [{"command": "add", "parent": "/", "type": "slide", "props": {"layout": "blank"}},
             {"command": "add", "parent": "/slide[2]", "type": "shape", "props": {"text": "Missing slide"}}],
        ):
            response = client.post("/v1/officecli/presentations/create", headers=headers,
                json={"output_name": "invalid.pptx", "commands": commands})
            self.assertEqual(response.status_code, 502, response.text)
            self.assertIn("officecli", response.json()["detail"])
            self.assertIsNone(files.result)
            self.assertEqual(files.attached, [])


if __name__ == "__main__":
    unittest.main()
