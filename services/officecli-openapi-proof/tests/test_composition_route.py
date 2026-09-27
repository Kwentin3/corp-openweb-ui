from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from officecli_openapi_proof.app import create_app
from officecli_openapi_proof.openwebui_client import NativeAttachment, OpenWebUiFailure, OpenWebUiUnauthorized
from test_bounded_inspection import HEADERS
from test_guidance_routes import RecordingOfficeCli, RecordingOpenWebUi, settings


@pytest.fixture
def native_files(tmp_path, monkeypatch):
    # The child is the real installed-module route; pytest's pythonpath is not
    # inherited by Python subprocesses, so expose this checkout explicitly.
    monkeypatch.setenv("PYTHONPATH", str(Path(__file__).resolve().parents[1]))
    paths = {}
    for index, name in enumerate(("source-a", "source-b")):
        path = tmp_path / f"{name}.xlsx"
        w = Workbook()
        w.active["A1"] = index + 3
        w.active["B1"] = "=A1*2"
        w.save(path)
        paths[name] = path

    class Files(RecordingOpenWebUi):
        def resolve_nearest_xlsx_attachments(self, *args):
            return [NativeAttachment(k, p.name) for k, p in paths.items()]

        def download(self, file_id, authorization, destination):
            self.calls.append(("download", file_id))
            destination.write_bytes(paths[file_id].read_bytes())

    return Files(), paths


def payload():
    return {"output_name": "combined.xlsx", "sources": [
        {"file_id": "source-a", "target_sheet": "Jan"},
        {"file_id": "source-b", "target_sheet": "Feb"},
    ]}


def test_composition_runs_real_worker_and_publishes_verified_native_file(native_files, tmp_path):
    files, paths = native_files
    before = {k: p.read_bytes() for k, p in paths.items()}
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=payload())
    assert response.status_code == 200, response.text
    assert response.json()["receipt"]["verified"]["formulas"] == 2
    assert response.json()["result_file_id"] == "result-file-id"
    artifact = tmp_path / "published.xlsx"
    artifact.write_bytes(files.uploaded_bytes)
    w = load_workbook(artifact)
    assert w.sheetnames == ["Jan", "Feb"]
    assert w["Jan"]["B2"].value == "='Jan'!A2*2"
    assert w["Feb"]["A2"].value == 4
    assert {k: p.read_bytes() for k, p in paths.items()} == before
    assert [c[0] for c in files.calls].count("upload") == 1
    assert [c[0] for c in files.calls].count("attach") == 1
    assert response.json()["download_url"] == f"/api/v1/files/{response.json()['result_file_id']}/content"
    assert len(response.content) < 3000


def test_omitted_source_cannot_be_reported_as_complete(native_files):
    files, _ = native_files
    request = payload()
    request["sources"].pop()
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=request)
    assert response.status_code == 422
    assert response.json()["detail"]["missing_file_ids"] == ["source-b"]
    assert files.calls == []


def test_unavailable_formula_dependency_prevents_publication(native_files):
    files, paths = native_files
    w = load_workbook(paths["source-a"])
    w.active["B1"] = "=Missing!A1"
    w.save(paths["source-a"])
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=payload())
    assert response.status_code == 422
    assert not files.uploaded


def test_failed_native_attachment_removes_orphan_upload(native_files):
    files, _ = native_files
    def unavailable(*args):
        raise OpenWebUiFailure("attachment denied")
    files.attach = unavailable
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=payload())
    assert response.status_code == 502
    assert ("delete", "result-file-id") in files.calls


def test_native_authorization_failure_prevents_work(native_files):
    files, _ = native_files
    def denied(*args):
        raise OpenWebUiUnauthorized("denied")
    files.verify_session = denied
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=payload())
    assert response.status_code == 401
    assert files.calls == []


def test_changed_native_source_prevents_publication(native_files):
    files, paths = native_files
    original_download = files.download
    downloads = 0
    def changing_source(file_id, authorization, destination):
        nonlocal downloads
        downloads += 1
        original_download(file_id, authorization, destination)
        if downloads == 3:
            destination.write_bytes(paths["source-b"].read_bytes())
    files.download = changing_source
    response = TestClient(create_app(RecordingOfficeCli(), files, settings())).post(
        "/v1/officecli/spreadsheets/compose", headers=HEADERS, json=payload())
    assert response.status_code == 502
    assert not files.uploaded


@pytest.mark.parametrize("new_upload", [False, True])
def test_followup_completeness_uses_user_uploads_on_active_branch(monkeypatch, new_upload):
    import httpx
    from officecli_openapi_proof.openwebui_client import HttpOpenWebUiClient

    messages = {
        "upload": {"role": "user", "parentId": None, "files": [
            {"id": "a", "name": "Jan.xlsx"}, {"id": "b", "name": "Feb.xlsx"}]},
        "result": {"role": "assistant", "parentId": "upload", "files": [
            {"id": "generated", "name": "combined.xlsx"}]},
        "followup": {"role": "user", "parentId": "result", "files":
            [{"id": "replacement", "name": "new.xlsx"}] if new_upload else []},
        "now": {"role": "assistant", "parentId": "followup", "files": []},
        "other-branch": {"role": "user", "parentId": "result", "files": [
            {"id": "unrelated", "name": "other.xlsx"}]},
    }

    def request(method, url, **kwargs):
        assert method == "GET" and url.endswith("/api/v1/chats/chat")
        assert kwargs["headers"]["Authorization"] == "Bearer caller"
        return httpx.Response(200, json={"chat": {"history": {"messages": messages}}},
                              request=httpx.Request(method, url))

    monkeypatch.setattr("officecli_openapi_proof.openwebui_client.httpx.request", request)
    client = HttpOpenWebUiClient("http://openwebui:8080", 30)
    sources = client.resolve_nearest_xlsx_attachments("chat", "now", "Bearer caller")
    assert [f.file_id for f in sources] == (["replacement"] if new_upload else ["a", "b"])
    # Editing still resolves the nearest generated file when no new upload exists.
    assert client.resolve_nearest_xlsx_attachment("chat", "now", "Bearer caller") == (
        "replacement" if new_upload else "generated")
