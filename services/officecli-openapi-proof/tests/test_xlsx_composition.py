from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
import pytest

from officecli_openapi_proof.xlsx_composition import (
    CompositionUnsupported, Source, compose, relocate_formula,
)


def test_relocation_preserves_absolute_references_and_cross_sheet_dependencies():
    assert relocate_formula("=SUM(A1:$B$2)+'Day 2'!$C3", "Day 1", "Jan 26", {"Day 1": 1, "Day 2": 9}, 0) == "=SUM('Jan 26'!A2:$B$3)+'Jan 26'!$C12"
    assert relocate_formula('="A1"&A1', "Day 1", "Jan 26", {"Day 1": 1}, 0) == '''="A1"&'Jan 26'!A2'''
    assert relocate_formula("='[1]Old Sheet'!$A$1", "Day 1", "Jan", {"Day 1": 1}, 3) == "='[4]Old Sheet'!$A$1"


@pytest.mark.parametrize("formula", ["=INDIRECT(\"A1\")", "=ROW(A1)", "=SUM(A:A)", "=Missing!A1", "=SomeName"])
def test_unsupported_semantics_fail_visibly(formula):
    with pytest.raises(CompositionUnsupported):
        relocate_formula(formula, "Day", "Jan", {"Day": 1}, 0)


def test_source_backed_composition_preserves_formulas_values_and_original_bytes(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    w = Workbook()
    w.active.title = "Day 1"
    w.active.append([3, "=A1*2", "= 'Day 2'!$A$1"])
    w.create_sheet("Day 2")["A1"] = 7
    w["Day 2"]["A2"] = "=SUM('Day 1'!A1:B1)"
    w.save(a)
    w = Workbook()
    w.active["A1"] = "a literal =SUM(A1)"
    w.save(b)
    originals = [p.read_bytes() for p in (a, b)]
    result = tmp_path / "result.xlsx"
    receipt = compose([Source(a, "Jan"), Source(b, "Feb")], result)
    r = load_workbook(result)
    assert r.sheetnames == ["Jan", "Feb"]
    assert r["Jan"]["A2"].value == 3
    assert r["Jan"]["B2"].value == "='Jan'!A2*2"
    assert r["Jan"]["C2"].value == "= 'Jan'!$A$5"
    assert r["Jan"]["A5"].value == 7
    assert r["Jan"]["A6"].value == "=SUM('Jan'!A2:B2)"
    assert r["Feb"]["A2"].value == "a literal =SUM(A1)"
    assert receipt["source_workbooks"] == 2
    assert receipt["source_sheets"] == 3
    assert receipt["formulas"] == 3
    assert [p.read_bytes() for p in (a, b)] == originals


def test_formula_cache_survives_without_becoming_a_constant(tmp_path):
    source = tmp_path / "source.xlsx"
    staged = tmp_path / "cached-source.xlsx"
    w = Workbook()
    w.active["A1"] = 3
    w.active["B1"] = "=A1*2"
    w.save(source)
    with ZipFile(source) as zin, ZipFile(staged, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                root.find('.//s:c[@r="B1"]/s:v', ns).text = "6"
                data = ET.tostring(root)
            zout.writestr(item, data)
    result = tmp_path / "result.xlsx"
    compose([Source(staged, "Month")], result)
    assert load_workbook(result)["Month"]["B2"].value == "='Month'!A2*2"
    assert load_workbook(result, data_only=True)["Month"]["B2"].value == 6


def test_blank_referenced_rows_do_not_become_next_sheet_values(tmp_path):
    source = tmp_path / "source.xlsx"
    w = Workbook()
    w.active.title = "First"
    w.active["A1"] = "=A10"
    w.create_sheet("Second")["A1"] = 99
    w.save(source)
    result = tmp_path / "result.xlsx"
    compose([Source(source, "Month")], result)
    r = load_workbook(result)["Month"]
    assert r["A2"].value == "='Month'!A11"
    assert r["A11"].value is None
    assert r["A14"].value == 99


@pytest.mark.parametrize("changed", [999, "=42", "=A3", "='Other'!A2*2"])
def test_independent_verification_rejects_changed_result_artifact(tmp_path, changed):
    from officecli_openapi_proof.xlsx_verification import verify, CompositionVerificationError
    source, result = tmp_path / "source.xlsx", tmp_path / "result.xlsx"
    w = Workbook()
    w.active["A1"] = 3
    w.active["B1"] = "=A1*2"
    w.save(source)
    compose([Source(source, "Month")], result)
    assert verify([source], result)["formulas"] == 1
    altered = load_workbook(result)
    altered["Month"]["B2"] = changed
    altered.save(result)
    with pytest.raises(CompositionVerificationError):
        verify([source], result)
