"""Validate native downloaded canaries against the committed source workbooks."""
import base64
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

out = Path(__file__).parent
source = out.parent / "officecli-multi-xlsx"
results = []
for receipt_path in map(Path, sys.argv[1:]):
    receipt = json.loads(receipt_path.read_text(encoding="utf-8-sig"))
    answer = receipt["messages"][-1]
    assert answer["done"] is True and not answer.get("error")
    assert len(answer["files"]) == 1
    file = answer["files"][0]
    assert file["status"] == file["metadata_status"] == 200
    assert file["user_id"] == receipt["user_id"]
    target = out / file["name"]
    encoded = file.pop("base64", None)
    payload = base64.b64decode(encoded) if encoded is not None else target.read_bytes()
    target.write_bytes(payload)
    wb = load_workbook(target)
    assert wb.sheetnames == ["Янв26", "Фев26"]
    expected_totals = [350, 550]
    for index, (name, sheet) in enumerate(zip(["jan.xlsx", "feb.xlsx"], wb)):
        values = list(load_workbook(source / name).active.values)
        assert (sheet.max_row, sheet.max_column) == (4, 4)
        assert list(sheet.values) == [
            (*values[0], "Amount"),
            (*values[1], "=B2*C2"),
            (*values[2], "=B3*C3"),
            ("Итого", None, None, "=SUM(D2:D3)"),
        ]
        total = sum(sheet.cell(row, 2).value * sheet.cell(row, 3).value for row in [2, 3])
        assert total == expected_totals[index]
    inspected = {c["file_id"] for c in answer["calls"] if c["name"] == "inspect_office_spreadsheet"}
    inputs = {f["id"] for m in receipt["messages"] if m["role"] == "user" for f in m["files"]}
    assert inspected == inputs and len(inputs) == 2
    assert any(c.get("result_file_id") == file["id"] for c in answer["calls"])
    # Preserve native ownership/tool receipts, not fixture RAG text or inline binary.
    for m in receipt["messages"]:
        m.pop("sources", None)
    (out / receipt_path.name).write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    results.append({"model": answer["model"], "role": "admin", "receipt": receipt_path.name,
                    "workbook": target.name, "sha256": hashlib.sha256(payload).hexdigest(),
                    "source_cells_formulas_topology_totals": "PASS", "totals": expected_totals})
(out / "acceptance.json").write_text(json.dumps({"results": results}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(results, ensure_ascii=False, indent=2))
