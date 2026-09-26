"""Verify downloaded product output against both synthetic source workbooks.

Usage: python verify_workbook.py path/to/output.xlsx
Requires openpyxl; checks source values, exact topology, formulas and cached totals.
"""
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def verify(path: Path) -> dict:
    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    assert formulas.sheetnames == ["Янв26", "Фев26"], formulas.sheetnames
    sources = Path(__file__).parent
    for name, source, total in (("Янв26", "jan.xlsx", 350), ("Фев26", "feb.xlsx", 550)):
        sheet = formulas[name]
        original = load_workbook(sources / source).active
        assert (sheet.max_row, sheet.max_column) == (4, 4)
        assert tuple(sheet.cell(1, c).value for c in range(1, 5)) == (
            "Product", "Quantity", "Price", "Amount"
        )
        for row in (2, 3):
            assert tuple(sheet.cell(row, c).value for c in range(1, 4)) == tuple(
                original.cell(row, c).value for c in range(1, 4)
            )
            assert sheet.cell(row, 4).value == f"=B{row}*C{row}"
            assert cached[name].cell(row, 4).value == original.cell(row, 2).value * original.cell(row, 3).value
        assert sheet["A4"].value == "Итого"
        assert sheet["B4"].value is None and sheet["C4"].value is None
        assert sheet["D4"].value == "=SUM(D2:D3)"
        assert cached[name]["D4"].value == total
        errors = [cell.coordinate for row in cached[name] for cell in row if cell.data_type == "e"]
        assert not errors, errors
    return {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "sheets": formulas.sheetnames, "totals": [350, 550], "passed": True}


if __name__ == "__main__":
    print(json.dumps(verify(Path(sys.argv[1])), ensure_ascii=True))
